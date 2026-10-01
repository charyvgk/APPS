#!/usr/bin/env python3
"""Build the leadership briefing as a Word document."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "docs" / "presentation" / "figures"
OUT = ROOT / "docs" / "presentation" / "Application-Architecture-Briefing.docx"

NAVY = "0F2C4C"
TEAL = "1F6F6A"
GOLD = "8C6A2F"
INK = "243038"
MUTED = "5C6670"
PAPER = "F6F4EF"
CARD = "EEF3F6"
SECURITY = "6E2B3A"
MONITOR = "1B4F72"
WHITE = "FFFFFF"
RULE = "D9D3C7"

FONT_REG = "/usr/share/fonts/truetype/macos/Inter-Regular.ttf"
FONT_MED = "/usr/share/fonts/truetype/macos/Inter-Medium.ttf"
FONT_SEM = "/usr/share/fonts/truetype/macos/Inter-SemiBold.ttf"
FONT_BOLD = "/usr/share/fonts/truetype/macos/Inter-Bold.ttf"


def rgb(hex_color):
    value = hex_color.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def shade(cell, color):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    existing = tc_pr.find(qn("w:shd"))
    if existing is not None:
        tc_pr.remove(existing)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def set_borders(cell, color="FFFFFF", size="4"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    existing = tc_pr.find(qn("w:tcBorders"))
    if existing is not None:
        tc_pr.remove(existing)
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)
        borders.append(element)
    tc_pr.append(borders)


def margins(cell, top=60, bottom=60, left=80, right=80):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is not None:
        tc_pr.remove(tc_mar)
    tc_mar = OxmlElement("w:tcMar")
    for name, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)
    tc_pr.append(tc_mar)


def prevent_row_split(row):
    tr = row._tr
    tr_pr = tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    tr_pr.append(cant)


def set_run_font(run, name="Inter", size=11, bold=False, color=INK, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = rgb(color)


def add_text(paragraph, text, size=11, bold=False, color=INK, italic=False):
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color, italic=italic)
    return run


def style_paragraph(paragraph, before=0, after=8, align="left", line=1.08):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    paragraph.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }[align]


def h1(doc, text):
    paragraph = doc.add_paragraph()
    style_paragraph(paragraph, before=0, after=6)
    add_text(paragraph, text, size=22, bold=True, color=NAVY)
    return paragraph


def h2(doc, text, color=NAVY):
    paragraph = doc.add_paragraph()
    style_paragraph(paragraph, before=10, after=4)
    add_text(paragraph, text, size=16, bold=True, color=color)
    return paragraph


def h3(doc, text, color=TEAL):
    paragraph = doc.add_paragraph()
    style_paragraph(paragraph, before=8, after=2)
    add_text(paragraph, text, size=13, bold=True, color=color)
    return paragraph


def body(doc, text, size=11, after=8, color=INK, bold=False):
    paragraph = doc.add_paragraph()
    style_paragraph(paragraph, after=after)
    add_text(paragraph, text, size=size, color=color, bold=bold)
    return paragraph


def bullet(doc, text, level=0):
    paragraph = doc.add_paragraph()
    style_paragraph(paragraph, after=3)
    paragraph.paragraph_format.left_indent = Cm(0.75 + level * 0.5)
    paragraph.paragraph_format.first_line_indent = Cm(-0.35)
    add_text(paragraph, "•  " + text, size=11)
    return paragraph


def caption(doc, text):
    paragraph = doc.add_paragraph()
    style_paragraph(paragraph, before=2, after=10)
    add_text(paragraph, text, size=9, italic=True, color=MUTED)
    return paragraph


def page_break(doc):
    doc.add_page_break()


def set_col_widths(table, widths):
    table.autofit = False
    table.allow_autofit = False
    for row in table.rows:
        for index, width in enumerate(widths):
            row.cells[index].width = Inches(width)


def cell_para(cell, text, size=10.5, bold=False, color=INK, align="left", after=2, before=0):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    style_paragraph(paragraph, before=before, after=after, align=align)
    add_text(paragraph, text, size=size, bold=bold, color=color)
    return paragraph


def add_cell_para(cell, text, size=10.5, bold=False, color=INK, after=2):
    paragraph = cell.add_paragraph()
    style_paragraph(paragraph, after=after)
    add_text(paragraph, text, size=size, bold=bold, color=color)
    return paragraph


def simple_table(doc, headers, rows, widths, header_color=NAVY):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        shade(cell, header_color)
        set_borders(cell, "FFFFFF", "4")
        margins(cell, 50, 50, 70, 70)
        cell_para(cell, header, size=10, bold=True, color=WHITE)
    for row_index, row in enumerate(rows):
        for index, value in enumerate(row):
            cell = table.rows[row_index + 1].cells[index]
            shade(cell, WHITE if row_index % 2 == 0 else PAPER)
            set_borders(cell, RULE, "4")
            margins(cell, 46, 46, 70, 70)
            cell_para(cell, value, size=10)
        prevent_row_split(table.rows[row_index + 1])
    prevent_row_split(table.rows[0])
    set_col_widths(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return table


def component_card(doc, name, role, why, preferred, security, monitoring):
    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    top = table.rows[0].cells[0].merge(table.rows[0].cells[1])
    shade(top, NAVY)
    set_borders(top, NAVY, "4")
    margins(top, 50, 40, 90, 90)
    cell_para(top, name, size=13, bold=True, color=WHITE, after=0)

    role_cell = table.rows[1].cells[0].merge(table.rows[1].cells[1])
    shade(role_cell, "E7EEF2")
    set_borders(role_cell, NAVY, "4")
    margins(role_cell, 40, 40, 90, 90)
    cell_para(role_cell, role, size=10.5, color=INK, after=0)

    pairs = [
        (2, "Why it is in the architecture", why, "Why we preferred it", preferred, TEAL),
        (3, "What Security gains", security, "What Monitoring gains", monitoring, MONITOR),
    ]
    for row_index, left_label, left_text, right_label, right_text, accent in pairs:
        left = table.rows[row_index].cells[0]
        right = table.rows[row_index].cells[1]
        for cell, label, text in ((left, left_label, left_text), (right, right_label, right_text)):
            shade(cell, WHITE)
            set_borders(cell, RULE, "4")
            margins(cell, 50, 50, 80, 80)
            cell_para(cell, label, size=9, bold=True, color=accent, after=1)
            add_cell_para(cell, text, size=10, after=0)
        prevent_row_split(table.rows[row_index])
    prevent_row_split(table.rows[0])
    prevent_row_split(table.rows[1])
    set_col_widths(table, [3.45, 3.45])
    spacer = doc.add_paragraph()
    style_paragraph(spacer, after=8)
    return table


def add_page_field(paragraph):
    run = paragraph.add_run()
    set_run_font(run, size=9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def configure_section(section):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.6)
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.6)
    section.header_distance = Cm(0.6)
    section.footer_distance = Cm(0.5)
    section.different_first_page_header_footer = True

    header = section.header
    header.is_linked_to_previous = False
    paragraph = header.paragraphs[0]
    style_paragraph(paragraph, after=2)
    add_text(paragraph, "Application architecture briefing", size=9, color=NAVY, bold=True)
    add_text(paragraph, "   ·   Azure, hybrid, and on-premises", size=9, color=MUTED)
    rule = header.add_paragraph()
    style_paragraph(rule, before=1, after=0)
    p_pr = rule._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), NAVY)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)

    footer = section.footer
    footer.is_linked_to_previous = False
    footer_para = footer.paragraphs[0]
    style_paragraph(footer_para, after=0)
    add_text(footer_para, "Security   ·   Monitoring   ·   Architecture          ", size=9, color=MUTED)
    add_page_field(footer_para)

    first_footer = section.first_page_footer
    first_footer.is_linked_to_previous = False
    first_footer.paragraphs[0].text = ""


def font(path, size):
    return ImageFont.truetype(path, size)


def wrap_text(draw, text, face, max_width):
    lines = []
    for paragraph in text.split("\n"):
        current = ""
        for word in paragraph.split():
            trial = word if not current else f"{current} {word}"
            if draw.textlength(trial, font=face) <= max_width:
                current = trial
            else:
                if current:
                    lines.append(current)
                current = word
        lines.append(current)
    return lines


def draw_lines(draw, lines, xy, face, fill, line_gap=4):
    x, y = xy
    ascent, descent = face.getmetrics()
    height = ascent + descent
    for line in lines:
        draw.text((x, y), line, font=face, fill=fill)
        y += height + line_gap
    return y


def card(draw, box, fill, title, body_text, title_fill, body_fill, title_size=22, body_size=16):
    draw.rounded_rectangle(box, radius=18, fill=fill)
    x1, y1, x2, y2 = box
    title_font = font(FONT_SEM, title_size)
    body_font = font(FONT_REG, body_size)
    draw_lines(draw, wrap_text(draw, title, title_font, x2 - x1 - 36), (x1 + 18, y1 + 16), title_font, title_fill, 2)
    title_height = title_font.getmetrics()[0] + title_font.getmetrics()[1] + 8
    draw_lines(
        draw,
        wrap_text(draw, body_text, body_font, x2 - x1 - 36),
        (x1 + 18, y1 + 16 + title_height),
        body_font,
        body_fill,
        3,
    )


def arrow(draw, start, end, fill=(15, 44, 76)):
    draw.line([start, end], fill=fill, width=4)
    ex, ey = end
    sx, sy = start
    if abs(ex - sx) >= abs(ey - sy):
        direction = 1 if ex > sx else -1
        draw.polygon([(ex, ey), (ex - 14 * direction, ey - 8), (ex - 14 * direction, ey + 8)], fill=fill)
    else:
        direction = 1 if ey > sy else -1
        draw.polygon([(ex, ey), (ex - 8, ey - 14 * direction), (ex + 8, ey - 14 * direction)], fill=fill)


def save_architecture():
    image = Image.new("RGB", (1800, 980), "#F7F4EE")
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, 1800, 78), fill="#0F2C4C")
    draw.text((36, 22), "Runtime architecture", font=font(FONT_BOLD, 32), fill="white")
    draw.text((430, 28), "One entry, one identity, two places to run", font=font(FONT_REG, 22), fill="#D5E2EA")

    card(draw, (40, 120, 360, 560), "#0F2C4C", "Callers", "Remote users\nSite users\nPartners and branch systems", "white", "#E6EEF3", 26, 22)
    card(draw, (460, 120, 860, 330), "#1F6F6A", "Cloud gateway", "Front Door and API Management\nfor remote and partner traffic", "white", "#E7F3F1", 24, 20)
    card(draw, (460, 360, 860, 560), "#1F6F6A", "Site gateway", "Self-hosted API gateway\nfor people and systems on premises", "white", "#E7F3F1", 24, 20)
    card(draw, (960, 120, 1360, 560), "#1B4F72", "Services we own", "Customer\nCatalog\nOrders\nNotification worker", "white", "#E4EEF5", 26, 22)
    card(draw, (1460, 120, 1760, 560), "#6E2B3A", "Outside systems", "Entra ID\nOn-premises ERP\nPayment\nMail", "white", "#F8E9ED", 24, 20)

    arrow(draw, (360, 250), (460, 220))
    arrow(draw, (360, 430), (460, 460))
    arrow(draw, (860, 225), (960, 280))
    arrow(draw, (860, 460), (960, 400))
    arrow(draw, (1360, 340), (1460, 340))

    card(
        draw,
        (40, 620, 1760, 900),
        "#EFEAE1",
        "Same controls on Azure and on premises",
        "Entra ID issues every access token. Argo CD deploys the same signed image to AKS and to the on-premises cluster.\nAzure Arc registers the site for policy and identity. Key Vault holds secrets. Azure Monitor receives traces from both sites.\nA private network carries any call that must cross between Azure and the premises.",
        "#0F2C4C",
        "#243038",
        26,
        22,
    )
    image.save(FIG / "architecture.png", quality=95)


def save_security():
    image = Image.new("RGB", (1800, 860), "#F7F4EE")
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, 1800, 78), fill="#6E2B3A")
    draw.text((36, 22), "Security controls on the path", font=font(FONT_BOLD, 32), fill="white")

    steps = [
        (40, "1", "Edge", "WAF, TLS, quota,\nand request size"),
        (390, "2", "Token", "Signature, issuer,\naudience, scope"),
        (740, "3", "Private path", "No direct route\nto a service"),
        (1090, "4", "Service check", "Token checked\nagain, then ownership"),
        (1440, "5", "Secrets", "Vault access by\nworkload identity"),
    ]
    for x, number, title, text in steps:
        draw.rounded_rectangle((x, 140, x + 320, 430), radius=18, fill="#6E2B3A")
        draw.text((x + 22, 160), number, font=font(FONT_BOLD, 28), fill="#F2D3C8")
        draw.text((x + 70, 164), title, font=font(FONT_SEM, 26), fill="white")
        draw_lines(draw, text.split("\n"), (x + 22, 230), font(FONT_REG, 22), "white", 6)
    for x in (360, 710, 1060, 1410):
        arrow(draw, (x, 285), (x + 30, 285), fill=(110, 43, 58))

    bands = [
        (40, "Supply chain", "The pipeline signs the image. Both clusters refuse an unsigned image."),
        (620, "Change control", "Production syncs only after approval. Gateway policy is reviewed in Git."),
        (1200, "Audit", "Denied calls, privileged roles, and order writes are attributable."),
    ]
    for x, title, text in bands:
        draw.rounded_rectangle((x, 490, x + 540, 780), radius=18, fill="white", outline="#6E2B3A", width=3)
        draw.text((x + 24, 516), title, font=font(FONT_SEM, 24), fill="#6E2B3A")
        draw_lines(draw, wrap_text(draw, text, font(FONT_REG, 22), 490), (x + 24, 570), font(FONT_REG, 22), "#243038", 6)
    image.save(FIG / "security.png", quality=95)


def save_monitoring():
    image = Image.new("RGB", (1800, 860), "#F7F4EE")
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, 1800, 78), fill="#1B4F72")
    draw.text((36, 22), "Monitoring from both sites", font=font(FONT_BOLD, 32), fill="white")

    sources = [
        (40, "Gateway", "Every accept and reject\nwith a correlation id"),
        (400, "Services", "Traces, errors,\nand upstream calls"),
        (760, "Clusters", "Health, saturation,\nand deploy readiness"),
    ]
    for x, title, text in sources:
        draw.rounded_rectangle((x, 140, x + 320, 390), radius=18, fill="#1B4F72")
        draw.text((x + 22, 168), title, font=font(FONT_SEM, 26), fill="white")
        draw_lines(draw, text.split("\n"), (x + 22, 230), font(FONT_REG, 22), "white", 6)
        arrow(draw, (x + 160, 390), (x + 160, 470), fill=(27, 79, 114))

    draw.rounded_rectangle((40, 480, 1080, 760), radius=18, fill="#0F2C4C")
    draw.text((70, 510), "OpenTelemetry into Azure Monitor", font=font(FONT_SEM, 28), fill="white")
    draw_lines(
        draw,
        wrap_text(
            draw,
            "One workspace for Azure and the premises. Logs, metrics, and traces share the gateway correlation id. A site user and a remote user are diagnosed the same way.",
            font(FONT_REG, 22),
            980,
        ),
        (70, 570),
        font(FONT_REG, 22),
        "white",
        6,
    )

    uses = [(1140, "Operations", "Latency, errors, saturation"), (1140, "Security", "Denials and privileged use")]
    draw.rounded_rectangle((1140, 480, 1760, 600), radius=18, fill="#1F6F6A")
    draw.text((1170, 505), "Operations", font=font(FONT_SEM, 24), fill="white")
    draw.text((1170, 545), "Latency, errors, saturation", font=font(FONT_REG, 20), fill="white")
    draw.rounded_rectangle((1140, 630, 1760, 760), radius=18, fill="#6E2B3A")
    draw.text((1170, 655), "Security", font=font(FONT_SEM, 24), fill="white")
    draw.text((1170, 695), "Denials and privileged use", font=font(FONT_REG, 20), fill="white")
    image.save(FIG / "monitoring.png", quality=95)


def save_promotion():
    image = Image.new("RGB", (1800, 520), "#F7F4EE")
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, 1800, 78), fill="#0F2C4C")
    draw.text((36, 22), "Build once, then approve the promotion", font=font(FONT_BOLD, 30), fill="white")
    stages = [
        (40, "#1F6F6A", "Build", "Test, scan, sign\nOne image digest"),
        (390, "#1B4F72", "Dev and test", "Auto-sync\nReviewer for test"),
        (740, "#8C6A2F", "Staging", "Tech lead approval\nBoth sites"),
        (1090, "#6E2B3A", "Production", "Two approvals\nManual sync"),
        (1440, "#0F2C4C", "Gateway last", "Route published\nafter health"),
    ]
    for x, color, title, text in stages:
        draw.rounded_rectangle((x, 150, x + 320, 430), radius=18, fill=color)
        draw.text((x + 24, 185), title, font=font(FONT_SEM, 26), fill="white")
        draw_lines(draw, text.split("\n"), (x + 24, 250), font(FONT_REG, 22), "white", 8)
    for x in (360, 710, 1060, 1410):
        arrow(draw, (x, 290), (x + 30, 290))
    image.save(FIG / "promotion.png", quality=95)


COMPONENTS = [
    (
        "Azure Front Door and web application firewall",
        "The public front door for remote users and partners.",
        "It terminates TLS close to the caller, filters common web attacks, and presents one global address for the cloud gateway.",
        "We preferred a managed global entry with a firewall over publishing the API gateway directly to the internet. Site users do not travel this path. They use the gateway on the premises, so local traffic stays local.",
        "Attack traffic is reduced before it reaches application code. Firewall policy is a Security control, separate from business releases.",
        "Front Door provides the outside view of availability and latency, before the request enters the gateway.",
    ),
    (
        "Cloud API gateway",
        "Azure API Management enforces the API contract for callers who enter through Azure.",
        "It checks the access token, applies quota and size limits, selects the version, and hides every service behind one contract.",
        "We preferred a managed gateway over letting each service publish its own endpoint. One policy set is easier to audit. The same definition is published to the on-premises gateway, so the two sites do not drift apart in behavior.",
        "Authentication, quota, and schema rejection happen before a service is contacted. Denied calls are logged at one place.",
        "Every accept and reject carries a correlation id that Monitoring can follow through the services.",
    ),
    (
        "Self-hosted gateway on premises",
        "The same API Management rules, running on Azure Local or Azure Stack Hub.",
        "People and systems on the site call a local gateway. Policy still comes from the Azure control plane. Local services stay on the local network.",
        "We preferred this over sending every site call to Azure. Latency and data-placement expectations are met, while Security still has one policy source. A custom site proxy was set aside because it would create a second set of rules.",
        "The site has the same token rules as the cloud. A disconnected site keeps the last published policy and fails closed when signing keys can no longer be trusted.",
        "Site traffic is visible in the same workspace as cloud traffic, tagged by site.",
    ),
    (
        "Microsoft Entra ID",
        "The only issuer of access tokens for this application.",
        "People, partners, and services prove who they are to Entra ID. The gateway and each service trust that issuer and no other.",
        "We preferred Entra ID over API keys stored per service, and over a private user database inside the application. Revocation, conditional access, and tenant control stay with Security. Services do not invent their own login.",
        "One issuer, one tenant, and an allow list of client applications. Conditional access can be tightened without a software release.",
        "Sign-in failures and token rejections are distinguishable from application defects.",
    ),
    (
        "Active Directory and Entra Connect",
        "The existing site directory remains the source for employee accounts.",
        "Entra Connect publishes site accounts and groups into Entra ID. A federated domain can still check the password in Active Directory. The access token is issued by Entra ID.",
        "We preferred to use the directory the company already runs, rather than asking every site user to hold a second account. The application does not connect to Active Directory on each request.",
        "Account disablement in the site directory reaches the application through Entra ID. There is no shadow account list to drift.",
        "Directory sync health matters. A broken sync is an identity incident, not an application defect, and should be alarmed separately.",
    ),
    (
        "Downstream services",
        "Customer, Catalog, Orders, and the notification worker. These are the systems this program owns.",
        "Each service owns its data and one business capability. The gateway is the only public route to them. The same image runs in Azure and on premises.",
        "We preferred separate services over one large application so a catalog change has a small blast radius and a service can run next to the data it uses. We preferred the same image on both sites over maintaining two code lines.",
        "Each service checks the token again. A caller who bypasses the gateway still cannot act. Ownership rules, such as seeing only your own orders, live here.",
        "Each service reports liveness and readiness. Readiness includes its database, so Monitoring and deployment share one health signal.",
    ),
    (
        "Upstream systems",
        "Systems outside this application: the on-premises ERP, the payment provider, and mail.",
        "Orders is the only service that calls the ERP and payment. The notification worker is the only service that calls mail. Timeouts and retries are explicit.",
        "We preferred one client per outside system over allowing every service to integrate directly. Credentials, audit, and failure behavior stay in one place. A slow ERP cannot take down Catalog.",
        "Outbound credentials live in Key Vault. Egress is allow-listed. Payment and ERP calls are attributable to a user or a partner application.",
        "Upstream latency and errors are measured apart from our own errors, so an ERP incident is not reported as an application outage without evidence.",
    ),
    (
        "Argo CD",
        "The only tool that deploys this application onto a cluster.",
        "Git holds the approved image digest. Argo CD makes each cluster match that Git revision. Production does not move on a pipeline script. It moves on a manual sync after approval.",
        "We preferred Argo CD over manual deployment and over ad-hoc scripts. We also preferred it as the single reconciler. Azure Arc can run Flux, and we leave that extension off for these namespaces so two deployers cannot overwrite each other.",
        "Approvers see the exact change before production sync. Drift after that sync is returned to the approved revision. There is a record of who synced production.",
        "A failed sync or a failed readiness check is an alert. Deployment health is part of the monitoring view, not a side channel.",
    ),
    (
        "Azure Arc",
        "Registration of the on-premises cluster in Azure for inventory, policy, and identity.",
        "The site cluster is visible in the same tenant as AKS. Workload identity and Azure Policy can apply there. Arc does not deploy the application.",
        "We preferred a registered cluster over an unmanaged server room. Security and Monitoring can see the site without a separate inventory. We did not prefer Arc's Flux deployment for this program because Argo CD already owns release control.",
        "Policy compliance for the site is reported next to Azure. Identity for site workloads uses the same tenant.",
        "Arc closes the visibility gap. A site cluster that falls out of registration is an operational alert.",
    ),
    (
        "AKS and on-premises Kubernetes",
        "The two places the services run. AKS in Azure. Kubernetes on Azure Local or Azure Stack Hub.",
        "Both sites run the same kind of cluster, so probes, network policy, and admission rules mean the same thing everywhere.",
        "We preferred one runtime over App Service in Azure and virtual machines on premises. Two different runtimes would split Security baselines and Monitoring agents. Kubernetes is the common substrate that hybrid placement needs.",
        "Network policy and admission control are enforceable on both clusters. The service is not reachable except from the gateway and its allowed peers.",
        "The same agent model ships logs, metrics, and traces from both clusters.",
    ),
    (
        "Container Registry and the site cache",
        "One registry in Azure, and a connected registry or pull-through cache on premises.",
        "The pipeline pushes one signed image. Production on either site pulls that digest. The site cache lets the premises pull without a second build.",
        "We preferred build-once over building again on the site. A rebuild would mean production might not be the bits that were approved. A local cache is preferred over depending on a wide-open pull from Azure during a site release.",
        "Only the pipeline can publish the production repository. Pulls are authenticated. The digest is what promotion approves.",
        "Replication lag to the site cache is visible, so a promotion does not start on a site that does not yet have the image.",
    ),
    (
        "Image signing and admission",
        "The pipeline signs the image. Each cluster rejects an image that does not carry that signature.",
        "This binds what is running to the pipeline that tested and scanned it. It applies on AKS and on the on-premises cluster.",
        "We preferred admission that refuses unsigned images over a written rule that people should only deploy from the pipeline. A registry credential alone cannot prove the image was approved.",
        "This is the supply-chain control. A substituted or side-loaded image does not start. Security can audit the signature identity.",
        "An admission rejection is a deploy failure with a clear reason, surfaced in Argo CD and in the cluster log.",
    ),
    (
        "Key Vault and workload identity",
        "Secrets stay in Azure Key Vault. A workload receives a short-lived identity to read the secrets it needs.",
        "No password is stored in Git or in the image. The same pattern works for AKS and for an Arc-enabled cluster. Rotation does not require a new release.",
        "We preferred workload identity over long-lived connection strings in configuration. A secret in a file spreads, and rotation becomes a release. Vault access is limited to the workload that needs it.",
        "Security can revoke a workload, rotate a secret, and see who read it. Developers do not handle production passwords.",
        "Vault access failures show up as readiness or dependency errors, separate from user-facing defects.",
    ),
    (
        "Policy on the cluster and Azure Policy",
        "Kyverno or Gatekeeper enforces baseline rules on every cluster. Azure Policy reports Azure and the Arc site together.",
        "Required health probes, resource limits, and signed images are enforced at deploy time. Organizational policy is visible for both sites.",
        "We preferred enforcement in the cluster over a checklist in a wiki. We preferred one admission engine on both sites so a rule written for Azure is meaningful on premises. Azure Policy remains the executive compliance view.",
        "Exceptions are recorded. A release that skips probes or runs as an unrestricted account is refused.",
        "Policy failures are deploy events. They should page the releasing team, not appear a week later in an audit.",
    ),
    (
        "OpenTelemetry and Azure Monitor",
        "Services emit traces, metrics, and logs with the gateway correlation id. Both sites export to Azure Monitor.",
        "A request can be followed from the door, through Orders, to the ERP. Remote and site traffic use one workspace.",
        "We preferred one telemetry pipeline over a separate monitoring stack on premises. Two stacks would force the Monitoring department to reconcile incidents by hand. OpenTelemetry keeps the application independent of a single vendor agent inside the code.",
        "Security receives the audit stream: denied tokens, use of privileged roles, and sensitive business actions.",
        "This is the system of record for service health. Dashboards and alerts for both sites live here.",
    ),
    (
        "Queue for work that can finish later",
        "A durable queue between Orders and the notification worker.",
        "The caller receives an order id when the order is stored. Mail can retry without holding the user on the line.",
        "We preferred a queue over sending mail inside the order request. A mail outage should not fail a sale that has already been recorded. The queue also gives Monitoring a visible backlog.",
        "Messages do not carry raw credentials. The worker uses its own identity toward the mail provider.",
        "Queue age and dead letters are operational signals. A growing backlog is an early warning before users complain.",
    ),
    (
        "Private network between Azure and the site",
        "ExpressRoute or a site-to-site VPN for traffic that must cross.",
        "Service-to-service calls and the path from Azure to the on-premises ERP do not travel the public internet.",
        "We preferred a private path over public endpoints between the sites. The ERP remains on the site network. Callers still use only the gateway.",
        "The private path is an allow-listed route, not a flat network. Platform firewalls still separate the clusters.",
        "Link health is a dependency of hybrid calls. Monitoring treats the circuit as a component with its own alert.",
    ),
    (
        "APIOps for gateway configuration",
        "API routes, token policies, and backends are stored in Git and published to API Management after review.",
        "A gateway change uses the same pull request discipline as a service change. Publishing happens after the new services are healthy.",
        "We preferred reviewed configuration over editing the API portal by hand. Hand edits do not show in the promotion record and are easy to make on one site only.",
        "Token policy and quota changes are visible to Security before they are published. The previous revision remains available for rollback.",
        "A publish is an event. Monitoring can correlate an error spike with a gateway revision the same way it correlates one with an image digest.",
    ),
]


def build_document():
    FIG.mkdir(parents=True, exist_ok=True)
    save_architecture()
    save_security()
    save_monitoring()
    save_promotion()

    doc = Document()
    configure_section(doc.sections[0])
    normal = doc.styles["Normal"]
    normal.font.name = "Inter"
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Inter")

    # Cover
    banner = doc.add_table(rows=1, cols=1)
    cell = banner.cell(0, 0)
    shade(cell, NAVY)
    set_borders(cell, NAVY, "2")
    margins(cell, 220, 200, 180, 180)
    cell_para(cell, "ARCHITECTURE BRIEFING", size=12, bold=True, color="D5E2EA", after=8)
    add_cell_para(cell, "One application architecture for Azure, hybrid, and on-premises", size=26, bold=True, color=WHITE, after=8)
    add_cell_para(cell, "Prepared for Security, Monitoring, and department leadership", size=14, color="E6EEF3", after=4)
    add_cell_para(cell, "1 October 2026", size=12, color="D5E2EA", after=0)
    set_col_widths(banner, [6.9])

    body(doc, "This briefing explains the architecture we intend to build, why each component is there, and why we preferred it to the practical alternatives. Technical specifications remain in the architecture set. This document is the version to present.", after=8)
    audience = doc.add_table(rows=1, cols=3)
    labels = [
        (SECURITY, "Security", "Identity, edge, secrets, supply chain, and audit"),
        (MONITOR, "Monitoring", "One view of both sites, from the door to the ERP"),
        (TEAL, "Architecture", "The same process in Azure and on premises"),
    ]
    for index, (color, title, text) in enumerate(labels):
        cell = audience.cell(0, index)
        shade(cell, color)
        set_borders(cell, WHITE, "12")
        margins(cell, 80, 80, 80, 80)
        cell_para(cell, title, size=12, bold=True, color=WHITE, after=2)
        add_cell_para(cell, text, size=10, color=WHITE, after=0)
    set_col_widths(audience, [2.3, 2.3, 2.3])

    h2(doc, "What we are asking")
    body(doc, "Endorse the architecture and the component choices in this briefing. Security is asked to own identity, edge, secret, and supply-chain policy. Monitoring is asked to own the shared service view and the alert thresholds. Neither department is asked to accept a different design on premises.")

    page_break(doc)
    h1(doc, "Architecture")
    body(doc, "Callers reach a gateway and nothing behind it. The gateway checks the caller and routes to services this program owns. Those services call a small number of outside systems when the work leaves our boundary. Azure and the premises run the same services. They do not run a different process.")
    doc.add_picture(str(FIG / "architecture.png"), width=Inches(6.9))
    caption(doc, "Figure 1. Callers, gateways, services we own, and systems outside the application.")

    h2(doc, "Three placements, one process")
    simple_table(
        doc,
        ["Placement", "Who calls", "Where services run", "What stays the same"],
        [
            ["Azure", "Remote users and partners", "AKS", "Gateway policy, tokens, image, approvals"],
            ["On premises", "People and systems on the site", "Kubernetes on Azure Local or Azure Stack Hub", "The same policy, tokens, image, and approvals"],
            ["Hybrid", "Either caller, routed to the site that can serve them", "Both clusters", "A private link when a call must cross"],
        ],
        [1.3, 1.8, 2.0, 1.8],
    )
    body(doc, "A site user is served by the self-hosted gateway and the local services. A remote user is served by Front Door and the cloud gateway. A call crosses to the other site only when the data or the upstream system is there. Cross-site traffic uses ExpressRoute or a site-to-site VPN.")

    page_break(doc)
    h1(doc, "How a request is handled")
    bullet(doc, "The caller signs in with Entra ID and receives a short-lived access token.")
    bullet(doc, "The caller sends the request to the gateway for that site.")
    bullet(doc, "The gateway applies firewall, quota, and schema controls, then validates the token.")
    bullet(doc, "A failed check stops at the gateway. The service is not called.")
    bullet(doc, "The service validates the same token, applies ownership rules, and does the work.")
    bullet(doc, "If the work needs the ERP or a payment, only the Orders service makes that call.")
    bullet(doc, "Mail and other slow work are queued. The caller is not held for them.")
    body(doc, "The gateway is a control point, not a place for business rules. Whether a customer may buy a product is decided by Orders, because that decision needs the order record and the ERP. Whether the caller may attempt the call is decided by the gateway.")

    h2(doc, "Words used in this briefing")
    simple_table(
        doc,
        ["Term", "Meaning"],
        [
            ["Gateway", "The only supported door. Cloud gateway in Azure, self-hosted gateway on premises."],
            ["Downstream service", "A service this program owns and runs behind the gateway."],
            ["Upstream service", "A system outside the program. Inbound upstream calls us. Outbound upstream is called by us."],
            ["Site", "One cluster: AKS, or Kubernetes on Azure Local or Azure Stack Hub."],
            ["Promotion", "Moving the same image digest to the next environment after approval."],
        ],
        [1.7, 5.2],
    )

    page_break(doc)
    h1(doc, "Components, and why these ones")
    body(doc, "Each component below states what it is, why the architecture includes it, and why we preferred it to the alternative we would otherwise have to operate. The Security and Monitoring columns say what that department receives from the choice.")

    for name, role, why, preferred, security, monitoring in COMPONENTS:
        component_card(doc, name, role, why, preferred, security, monitoring)

    page_break(doc)
    h1(doc, "Security")
    body(doc, "Security’s outcome is a single control model. A caller is identified the same way in Azure and on premises. A service cannot be reached except through the gateway. A release cannot start unless it was built and signed by the pipeline. A production change has an approver.")
    doc.add_picture(str(FIG / "security.png"), width=Inches(6.9))
    caption(doc, "Figure 2. Controls from the edge through the service, plus supply chain, change, and audit.")

    h2(doc, "Identity and access", color=SECURITY)
    bullet(doc, "People use authorization code with PKCE. The application never sees the password.")
    bullet(doc, "Partners and branch systems use client credentials. They do not share a person’s login.")
    bullet(doc, "Workloads use workload identity to reach Key Vault. There is no long-lived password in Git.")
    bullet(doc, "Delegated scopes for people include Catalog.Read, Orders.Read, Orders.Write, and Customer.Read.")
    bullet(doc, "Application roles include Orders.Submit and Orders.ReadAll. Orders.ReadAll is limited to operations clients.")
    bullet(doc, "The gateway and the service both check signature, issuer, audience, tenant, expiry, and the required scope or role.")
    bullet(doc, "The service then checks ownership. A valid Orders.Read token still returns only that caller’s orders.")

    h2(doc, "What is deliberately refused", color=SECURITY)
    simple_table(
        doc,
        ["Practice set aside", "Reason"],
        [
            ["Passwords inside the application", "Entra ID and Active Directory already own them."],
            ["A separate user store on premises", "It would drift from the corporate directory."],
            ["API keys as the main application identity", "They are hard to scope and easy to copy."],
            ["Trusting an identity header without a token", "A header can be forged if a path around the gateway appears."],
            ["Editing production gateway policy by hand", "It would bypass review and the promotion record."],
            ["Unsigned or rebuilt-on-site images", "Production must be the digest that was approved."],
        ],
        [2.6, 4.3],
    )

    h2(doc, "Decisions requested from Security", color=SECURITY)
    simple_table(
        doc,
        ["Decision", "Recommendation"],
        [
            ["Entra ID is the only token issuer", "Endorse"],
            ["Token checks run at the gateway and again in the service", "Endorse"],
            ["Front Door web application firewall protects the cloud entry", "Endorse, and own the firewall policy"],
            ["Secrets live only in Key Vault and are read with workload identity", "Endorse"],
            ["Both clusters refuse unsigned images", "Endorse, and own the signing identity"],
            ["Privileged role Orders.ReadAll is tightly held", "Endorse the role list and the client allow list"],
            ["Production deploy requires the recorded approvals", "Endorse"],
        ],
        [4.6, 2.3],
    )

    page_break(doc)
    h1(doc, "Monitoring")
    body(doc, "Monitoring’s outcome is one operational picture. Azure and the premises report to Azure Monitor. A correlation id assigned at the gateway ties the caller’s request to the service trace and, when needed, to the ERP or payment call.")
    doc.add_picture(str(FIG / "monitoring.png"), width=Inches(6.9))
    caption(doc, "Figure 3. Signals from the gateway, the services, and the clusters land in one workspace.")

    h2(doc, "What is measured", color=MONITOR)
    simple_table(
        doc,
        ["Signal", "Where it starts", "Why leadership should see it"],
        [
            ["Availability and latency", "Gateway", "This is the caller’s experience, on each site."],
            ["Token and quota rejects", "Gateway", "Separates abuse or identity issues from defects."],
            ["Service errors", "Each downstream service", "Shows which capability is failing."],
            ["Upstream errors and latency", "Orders and the worker", "Shows an ERP, payment, or mail incident."],
            ["Readiness and sync result", "Argo CD and probes", "Shows whether a promotion is safe to expose."],
            ["Queue age and dead letters", "Notification queue", "Shows work accepted but not yet finished."],
            ["Private link health", "ExpressRoute or VPN", "Shows whether hybrid calls can complete."],
            ["Audit events", "Gateway and services", "Denied access, privileged reads, and order writes."],
        ],
        [1.8, 1.8, 3.3],
    )

    h2(doc, "Proposed starting targets", color=MONITOR)
    body(doc, "These targets are a starting point for Monitoring to confirm. They are measured at the gateway for the site that received the call. Caller errors, such as a missing token, do not count against availability.")
    simple_table(
        doc,
        ["Target", "Proposal"],
        [
            ["Order API availability", "99.9 percent of calls succeed in a month"],
            ["Catalog read latency on the local site", "95 percent complete within 300 milliseconds"],
            ["Gateway server errors", "Alert when they exceed 1 percent for five minutes"],
            ["Production readiness", "Alert when a production sync is not healthy"],
            ["Entra signing-key retrieval", "Alert when validation can no longer refresh keys"],
            ["Notification backlog", "Alert when queue age exceeds the agreed business window"],
        ],
        [2.8, 4.1],
    )
    body(doc, "We preferred one Azure Monitor workspace over a second monitoring product on premises. The department should not have to learn two consoles to answer “is the site healthy?” Security’s audit queries use the same correlation id, with a retention period Monitoring and Security agree together.")

    h2(doc, "Decisions requested from Monitoring", color=MONITOR)
    simple_table(
        doc,
        ["Decision", "Recommendation"],
        [
            ["Azure Monitor is the system of record for both sites", "Endorse"],
            ["Dashboards cover availability, latency, errors, saturation, and gateway rejects", "Own the dashboard set"],
            ["Alerts include failed production readiness, token-failure spikes, and ERP errors", "Own the paging route"],
            ["The proposed targets above", "Confirm or replace before go-live"],
            ["Audit retention", "Agree the period with Security"],
        ],
        [4.6, 2.3],
    )

    page_break(doc)
    h1(doc, "Release control")
    body(doc, "Speed comes from building the image once. Safety comes from refusing to rebuild it for production and from requiring approval before a higher environment moves. Both sites receive the same digest. A failed site does not force the other site to roll back.")
    doc.add_picture(str(FIG / "promotion.png"), width=Inches(6.9))
    caption(doc, "Figure 4. The digest is promoted. Production and the gateway route wait for approval and health.")

    simple_table(
        doc,
        ["Step", "Approval", "What moves"],
        [
            ["Build on main", "Pull request review", "A signed digest in the registry"],
            ["Dev", "None beyond the merge", "Both sites auto-sync"],
            ["Test", "One reviewer", "Both sites auto-sync"],
            ["Staging", "Tech lead", "Both sites"],
            ["Production", "Release manager and operations", "Manual Argo CD sync, per site"],
            ["Gateway route", "Published only after service health", "Cloud and site gateways"],
        ],
        [1.6, 2.6, 2.7],
    )
    body(doc, "Rollback returns Git to the previous digest and syncs again. There is no emergency rebuild. The previous gateway revision is still available and becomes current if the route must be withdrawn.")

    h2(doc, "Why this release model")
    bullet(doc, "Approvers look at a diff of the approved digest, not at a new build of unknown content.")
    bullet(doc, "Dev and test stay fast because they auto-sync. Staging and production stay gated.")
    bullet(doc, "Each site is a separate deployment, so an on-premises fault does not take Azure with it.")
    bullet(doc, "The gateway route is last. Callers are not sent to services that are not ready.")
    bullet(doc, "Feature flags let a build be deployed and kept dark until the business is ready to expose it.")

    page_break(doc)
    h1(doc, "Decisions to record")
    body(doc, "If the departments accept this briefing, the following choices become the architecture baseline. A later change is a new decision, not a local exception on one site.")
    simple_table(
        doc,
        ["Choice", "Baseline"],
        [
            ["Entry", "Gateway only. Front Door for remote callers. Self-hosted gateway for the site."],
            ["Identity", "Entra ID tokens. Site accounts arrive through Entra Connect."],
            ["Authorization", "Scopes and roles at the gateway and the service. Ownership inside the service."],
            ["Runtime", "The same image on AKS and on on-premises Kubernetes."],
            ["Deploy", "Argo CD. Arc is for inventory, policy, and identity, not a second deployer."],
            ["Promotion", "Build once. Approve to promote. Manual sync in production."],
            ["Secrets", "Key Vault and workload identity."],
            ["Supply chain", "Signed images. Admission on both clusters."],
            ["Observation", "OpenTelemetry to Azure Monitor from both sites."],
            ["Gateway change", "Reviewed in Git and published after service health."],
        ],
        [1.6, 5.3],
    )

    h2(doc, "What happens next")
    bullet(doc, "Security confirms the role list, the partner client list, firewall ownership, and signing identity.")
    bullet(doc, "Monitoring confirms the workspace, the alert routes, and the starting targets.")
    bullet(doc, "Architecture keeps the component set stable and writes any exception as a decision.")
    body(doc, "The supporting architecture set in this repository holds the interaction diagrams for placement, authentication, and Argo CD. This briefing is the document to present.", after=4)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_document()
